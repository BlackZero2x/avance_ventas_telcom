"""
asistencia_planilla.py — Informe diario de asistencia de vendedores PLANILLA.

Lee la hoja ASISTENCIA_BBDD vía API de Google Sheets (OAuth, mismo token.json
del proyecto) y la cruza con RRHH para mostrar por zonal y por supervisor:
  - Total vendedores PLANILLA
  - Asistentes del día
  - % Asistencia

Genera dos imágenes y las envía al grupo Canal Fija 2026 Gestión AUREN.

Uso:
    python asistencia_planilla/asistencia_planilla.py                          # hoy, al grupo
    python asistencia_planilla/asistencia_planilla.py --fecha 2026-05-26       # fecha específica
    python asistencia_planilla/asistencia_planilla.py --destino "51975155264@c.us"  # prueba a número
"""
import argparse
import logging
import os
import json
import sys
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote


# ── Cargar .env y proxy ───────────────────────────────────────────────────────

def _cargar_dotenv():
    env_path = Path(__file__).parent.parent / ".env"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        os.environ.setdefault(key.strip(), value.strip())


def _configurar_proxy():
    no_proxy = "localhost,127.0.0.1,.googleapis.com,.google.com,.gstatic.com"
    os.environ["NO_PROXY"] = no_proxy
    os.environ["no_proxy"] = no_proxy

    proxy_host = os.environ.get("HTTP_PROXY", "")
    user = os.environ.get("PROXY_USER", "")
    password = os.environ.get("PROXY_PASS", "")

    if proxy_host and user and password:
        encoded_user = quote(user, safe="")
        encoded_pass = quote(password, safe="")
        base = proxy_host.replace("http://", "").replace("https://", "")
        proxy_url = f"http://{encoded_user}:{encoded_pass}@{base}"
    elif proxy_host:
        proxy_url = proxy_host
    else:
        proxy_url = ""

    if proxy_url:
        os.environ["HTTP_PROXY"] = proxy_url
        os.environ["http_proxy"] = proxy_url
    else:
        for var in ("HTTP_PROXY", "http_proxy"):
            os.environ.pop(var, None)
    for var in ("HTTPS_PROXY", "https_proxy"):
        os.environ.pop(var, None)


_cargar_dotenv()
_configurar_proxy()

import pandas as pd
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

BASE_DIR = Path(__file__).parent.parent  # raíz del proyecto
sys.path.insert(0, str(BASE_DIR / "whatsapp_server"))
from wa_client import WhatsAppClient


# ── Constantes ────────────────────────────────────────────────────────────────

LOGS_DIR    = BASE_DIR / "logs"
TEMP_DIR    = BASE_DIR / "temp"
TOKEN_PATH  = BASE_DIR / "token.json"
CREDS_PATH  = BASE_DIR / "credentials.json"
CONFIG_PATH = BASE_DIR / "config.json"

LOGS_DIR.mkdir(exist_ok=True)
TEMP_DIR.mkdir(exist_ok=True)

SCOPES = [
    "https://www.googleapis.com/auth/gmail.modify",
    "https://www.googleapis.com/auth/drive",
    "https://www.googleapis.com/auth/spreadsheets",
]

# Sheet de asistencia — contiene tanto RRHH como ASISTENCIA_BBDD
ASISTENCIA_SHEET_ID = "1EGNnQYG51MROVf2tuxKmEqCkw9lfIXQqJBkJke3AITk"
RRHH_HOJA           = "RRHH"
ASISTENCIA_HOJA     = "ASISTENCIA_BBDD"

GRUPO_GESTION = "Canal Fija 2026 Gestión AUREN"

# Paleta de colores
COLOR_HEADER   = (17, 138, 178)   # #118AB2 — azul cielo (encabezados y totales)
COLOR_TOTAL    = (17, 138, 178)   # #118AB2 — igual que header
COLOR_FILA_PAR = (235, 247, 251)  # azul cielo muy claro para filas alternas
COLOR_VERDE    = (131, 235, 176)  # #83EBB0
COLOR_AMARILLO = (255, 232, 127)  # #FFE87F
COLOR_ROJO     = (247, 163, 183)  # #F7A3B7


# ── Logging ───────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler(
            LOGS_DIR / f"asistencia_{datetime.now().strftime('%Y%m%d')}.log",
            encoding="utf-8",
        ),
        logging.StreamHandler(sys.stdout),
    ],
)


# ── Autenticación Google ──────────────────────────────────────────────────────

def _autenticar_sheets():
    creds = None
    if TOKEN_PATH.exists():
        creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(str(CREDS_PATH), SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_PATH, "w") as f:
            f.write(creds.to_json())
    return build("sheets", "v4", credentials=creds)


def _leer_hoja_api(service, sheet_id: str, nombre_hoja: str) -> pd.DataFrame:
    resultado = service.spreadsheets().values().get(
        spreadsheetId=sheet_id,
        range=nombre_hoja,
    ).execute()
    valores = resultado.get("values", [])
    if not valores or len(valores) < 2:
        return pd.DataFrame()
    encabezados = valores[0]
    filas = valores[1:]
    filas_norm = [fila + [""] * (len(encabezados) - len(fila)) for fila in filas]
    return pd.DataFrame(filas_norm, columns=encabezados)


# ── Carga de datos ────────────────────────────────────────────────────────────

def cargar_rh(service) -> pd.DataFrame:
    logging.info(f"Cargando hoja {RRHH_HOJA} via API...")
    df = _leer_hoja_api(service, ASISTENCIA_SHEET_ID, RRHH_HOJA)
    if df.empty:
        logging.error(f"Hoja {RRHH_HOJA} vacia o no accesible.")
        return df
    df.columns = [c.strip().upper() for c in df.columns]
    logging.info(f"Columnas RRHH: {list(df.columns)}")

    if "ESQUEMA" in df.columns:
        df["ESQUEMA"] = df["ESQUEMA"].replace("PART-TIME", "PLANILLA")
        df = df[df["ESQUEMA"].str.upper() == "PLANILLA"].copy()
    else:
        logging.warning("RH sin columna ESQUEMA — usando todos los registros.")

    # Filtro 1: solo ACTIVOS
    if "ESTADO" in df.columns:
        antes = len(df)
        df = df[df["ESTADO"].str.upper().str.strip() == "ACTIVO"].copy()
        logging.info(f"Filtro ESTADO=ACTIVO: {antes} -> {len(df)}")
    else:
        logging.warning("RH sin columna ESTADO — no se aplica filtro de activos.")

    # Filtro 2: solo EN CAMPO
    if "FEEDBACK_RH" in df.columns:
        antes = len(df)
        df = df[df["FEEDBACK_RH"].str.upper().str.strip() == "EN CAMPO"].copy()
        logging.info(f"Filtro FEEDBACK_RH=EN CAMPO: {antes} -> {len(df)}")
    else:
        logging.warning("RH sin columna FEEDBACK_RH — no se aplica filtro de en campo.")

    # Normalizar columna de zona
    if "ZONA" in df.columns and "ZONAL" not in df.columns:
        df = df.rename(columns={"ZONA": "ZONAL"})

    df["DNI"] = pd.to_numeric(df.get("DNI", pd.Series(dtype="object")), errors="coerce")
    logging.info(f"Vendedores PLANILLA activos EN CAMPO: {len(df)}")
    return df


def cargar_asistencia(service, fecha_hoy: date) -> pd.DataFrame:
    logging.info(f"Cargando hoja {ASISTENCIA_HOJA} via API...")
    df = _leer_hoja_api(service, ASISTENCIA_SHEET_ID, ASISTENCIA_HOJA)
    df.columns = [c.strip().upper() for c in df.columns]
    logging.info(f"Columnas ASISTENCIA_BBDD: {list(df.columns)}")
    logging.info(f"Total filas crudas: {len(df)}")

    col_fecha = "DAY"
    if col_fecha not in df.columns:
        logging.error(f"No se encontro columna DAY en ASISTENCIA_BBDD. Columnas: {list(df.columns)}")
        return pd.DataFrame()

    df["_FECHA"] = pd.to_datetime(df[col_fecha], dayfirst=True, errors="coerce").dt.date
    df_hoy = df[df["_FECHA"] == fecha_hoy].copy()
    logging.info(f"Registros de asistencia para {fecha_hoy}: {len(df_hoy)}")

    df_hoy["DNI"] = pd.to_numeric(
        df_hoy.get("DNI", pd.Series(dtype="object")), errors="coerce"
    )
    return df_hoy


# ── Cálculo de tablas ─────────────────────────────────────────────────────────

def _detectar_columna(df: pd.DataFrame, candidatos: list) -> str | None:
    cols_upper = {c.upper(): c for c in df.columns}
    for c in candidatos:
        if c.upper() in cols_upper:
            return cols_upper[c.upper()]
    return None


def calcular_ausentes(df_rh: pd.DataFrame, df_asistencia: pd.DataFrame) -> str:
    """Retorna lista de ausentes organizada por ZONAL → SUPERVISOR → VENDEDOR."""
    dnis_asistentes = set(df_asistencia["DNI"].dropna().astype(int).tolist()) if not df_asistencia.empty else set()

    df = df_rh.copy()
    df["ASISTIO"] = df["DNI"].apply(
        lambda d: 1 if pd.notna(d) and int(d) in dnis_asistentes else 0
    )

    col_zonal = _detectar_columna(df, ["ZONAL", "ZONA", "REGION"])
    col_sup = _detectar_columna(df, ["SUPERVISOR", "SUP", "JEFE"])
    col_vend = _detectar_columna(df, ["VENDEDOR", "NOMBRE", "PERSON"])

    if not col_zonal or not col_vend:
        logging.warning("No se pudo detectar columnas para lista de ausentes")
        return ""

    ausentes = df[df["ASISTIO"] == 0]
    if ausentes.empty:
        return "✅ *Todos asistieron hoy*"

    lineas = ["⚠️ *Vendedores sin registro de asistencia:*", ""]
    zonal_actual = None
    sup_actual = None

    for _, row in ausentes.sort_values([col_zonal] + ([col_sup] if col_sup else [])).iterrows():
        zonal = row[col_zonal]
        sup = row[col_sup] if col_sup else None
        vend = row[col_vend]

        if zonal != zonal_actual:
            if zonal_actual is not None:
                lineas.append("")
            lineas.append(f"*{zonal}*")
            zonal_actual = zonal
            sup_actual = None

        if col_sup and sup != sup_actual:
            lineas.append(f"  {sup}")
            sup_actual = sup

        lineas.append(f"    • {vend}")

    return "\n".join(lineas)


def calcular_tablas(df_rh: pd.DataFrame, df_asistencia: pd.DataFrame):
    """Retorna (df_por_zonal, df_por_supervisor)."""
    dnis_asistentes = set(df_asistencia["DNI"].dropna().astype(int).tolist()) if not df_asistencia.empty else set()

    df = df_rh.copy()
    df["ASISTIO"] = df["DNI"].apply(
        lambda d: 1 if pd.notna(d) and int(d) in dnis_asistentes else 0
    )

    col_zonal = _detectar_columna(df, ["ZONAL", "ZONA", "REGION"])
    if col_zonal is None:
        logging.error("No se encontro columna ZONAL en RH.")
        return pd.DataFrame(), pd.DataFrame()

    # ── Por ZONAL ─────────────────────────────────────────────────────────────
    grp_z = df.groupby(col_zonal).agg(
        PLANILLA=("DNI", "count"),
        ASISTENTES=("ASISTIO", "sum"),
    ).reset_index().rename(columns={col_zonal: "ZONAL"})
    grp_z[["PLANILLA", "ASISTENTES"]] = grp_z[["PLANILLA", "ASISTENTES"]].astype(int)
    grp_z["%ASISTENCIA"] = grp_z.apply(
        lambda r: r["ASISTENTES"] / r["PLANILLA"] * 100 if r["PLANILLA"] > 0 else 0.0,
        axis=1,
    )
    grp_z = grp_z.sort_values("ZONAL").reset_index(drop=True)

    total_z = {
        "ZONAL": "TOTAL",
        "PLANILLA": int(grp_z["PLANILLA"].sum()),
        "ASISTENTES": int(grp_z["ASISTENTES"].sum()),
    }
    total_z["%ASISTENCIA"] = (
        total_z["ASISTENTES"] / total_z["PLANILLA"] * 100
        if total_z["PLANILLA"] > 0 else 0.0
    )
    df_zonal = pd.concat(
        [grp_z, pd.DataFrame([total_z])], ignore_index=True
    )[["ZONAL", "PLANILLA", "ASISTENTES", "%ASISTENCIA"]]

    # ── Por SUPERVISOR ────────────────────────────────────────────────────────
    col_sup = _detectar_columna(df, ["SUPERVISOR", "SUP", "JEFE"])
    if col_sup is None:
        logging.warning("Sin columna SUPERVISOR en RH — tabla supervisores omitida.")
        return df_zonal, pd.DataFrame()

    grp_s = df.groupby([col_zonal, col_sup]).agg(
        PLANILLA=("DNI", "count"),
        ASISTENTES=("ASISTIO", "sum"),
    ).reset_index().rename(columns={col_zonal: "ZONAL", col_sup: "SUPERVISOR"})
    grp_s[["PLANILLA", "ASISTENTES"]] = grp_s[["PLANILLA", "ASISTENTES"]].astype(int)
    grp_s["%ASISTENCIA"] = grp_s.apply(
        lambda r: r["ASISTENTES"] / r["PLANILLA"] * 100 if r["PLANILLA"] > 0 else 0.0,
        axis=1,
    )
    grp_s = grp_s.sort_values(["ZONAL", "SUPERVISOR"]).reset_index(drop=True)

    total_s = {
        "ZONAL": "TOTAL", "SUPERVISOR": "",
        "PLANILLA": int(grp_s["PLANILLA"].sum()),
        "ASISTENTES": int(grp_s["ASISTENTES"].sum()),
    }
    total_s["%ASISTENCIA"] = (
        total_s["ASISTENTES"] / total_s["PLANILLA"] * 100
        if total_s["PLANILLA"] > 0 else 0.0
    )
    df_supervisor = pd.concat(
        [grp_s, pd.DataFrame([total_s])], ignore_index=True
    )[["ZONAL", "SUPERVISOR", "PLANILLA", "ASISTENTES", "%ASISTENCIA"]]

    return df_zonal, df_supervisor


# ── Generación de imagen Excel ────────────────────────────────────────────────

def _rgb_fill(r, g, b):
    from openpyxl.styles import PatternFill
    return PatternFill(
        start_color=f"{r:02X}{g:02X}{b:02X}",
        end_color=f"{r:02X}{g:02X}{b:02X}",
        fill_type="solid",
    )


def _color_asistencia(pct: float):
    if pct >= 90:
        return COLOR_VERDE
    elif pct >= 70:
        return COLOR_AMARILLO
    return COLOR_ROJO


def _generar_xlsx(df: pd.DataFrame, titulo: str, ruta_xlsx: str):
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabla"
    n_cols = len(df.columns)
    borde = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    # Título
    ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
    c = ws["A1"]
    c.value = titulo
    c.font = Font(name="Aptos Narrow", bold=True, size=13, color="FFFFFF")
    c.fill = _rgb_fill(*COLOR_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Encabezados
    for ci, col_nombre in enumerate(df.columns, start=1):
        c = ws.cell(row=2, column=ci, value=col_nombre)
        c.font = Font(name="Aptos Narrow", bold=True, size=10, color="FFFFFF")
        c.fill = _rgb_fill(*COLOR_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde
    ws.row_dimensions[2].height = 18

    # Filas de datos
    n_filas_datos = len(df)
    for fi, (_, row) in enumerate(df.iterrows(), start=3):
        es_total = (fi - 3) == n_filas_datos - 1
        fill_fila = _rgb_fill(*COLOR_TOTAL) if es_total else (
            _rgb_fill(*COLOR_FILA_PAR) if fi % 2 == 0 else None
        )
        for ci, col_nombre in enumerate(df.columns, start=1):
            valor = row[col_nombre]
            c = ws.cell(row=fi, column=ci)
            if col_nombre == "%ASISTENCIA":
                c.value = valor / 100 if isinstance(valor, (int, float)) else 0
                c.number_format = "0.0%"
                r, g, b = _color_asistencia(float(valor))
                c.fill = _rgb_fill(r, g, b)
                c.font = Font(name="Aptos Narrow", bold=True, size=9, color="000000")
            else:
                c.value = int(valor) if isinstance(valor, float) and valor == int(valor) else valor
                c.font = Font(name="Aptos Narrow", bold=es_total, size=9,
                              color="FFFFFF" if es_total else "000000")
                if fill_fila:
                    c.fill = fill_fila
            c.alignment = Alignment(
                horizontal="left" if col_nombre in ("ZONAL", "SUPERVISOR") else "center",
                vertical="center",
            )
            c.border = borde
        ws.row_dimensions[fi].height = 15

    # Anchos de columna
    anchos_fijos = {"PLANILLA": 11, "ASISTENTES": 12, "%ASISTENCIA": 13}
    for ci, col_nombre in enumerate(df.columns, start=1):
        col_letter = get_column_letter(ci)
        if col_nombre in ("ZONAL", "SUPERVISOR"):
            max_len = max(
                (len(str(v)) for v in df[col_nombre] if v is not None),
                default=len(col_nombre),
            )
            ws.column_dimensions[col_letter].width = min(max(max_len, len(col_nombre)) + 2, 40)
        else:
            ws.column_dimensions[col_letter].width = anchos_fijos.get(col_nombre, 10)

    wb.save(ruta_xlsx)
    logging.info(f"Excel temporal guardado: {ruta_xlsx}")


def _capturar_xlsx_a_png(ws_xw, n_filas: int, n_cols: int, ruta_png: str) -> bool:
    from openpyxl.utils import get_column_letter
    from PIL import ImageGrab, Image
    import win32gui
    import ctypes

    ultima_col = get_column_letter(n_cols)
    xl_range = ws_xw.range(f"A1:{ultima_col}{n_filas}")

    try:
        app = ws_xw.book.app
        app.api.Visible = True
        hwnd = [0]
        def _cb(h, _):
            if win32gui.IsWindowVisible(h) and "Microsoft Excel" in win32gui.GetWindowText(h):
                hwnd[0] = h
                return False
            return True
        win32gui.EnumWindows(_cb, None)
        if hwnd[0]:
            ctypes.windll.user32.AllowSetForegroundWindow(
                ctypes.windll.kernel32.GetCurrentProcessId()
            )
            win32gui.ShowWindow(hwnd[0], 9)
            win32gui.SetForegroundWindow(hwnd[0])
        time.sleep(1)
    except Exception:
        pass

    for intento in range(3):
        try:
            xl_range.api.CopyPicture(Appearance=1, Format=2)
            time.sleep(1.5)
            img = ImageGrab.grabclipboard()
            if img:
                nuevo_ancho = int(img.width * 2.5)
                nuevo_alto  = int(img.height * 2.5)
                img = img.resize((nuevo_ancho, nuevo_alto), resample=Image.LANCZOS)
                img.save(ruta_png, "PNG")
                logging.info(f"Imagen guardada: {ruta_png} ({nuevo_ancho}x{nuevo_alto})")
                return True
        except Exception as e:
            logging.warning(f"CopyPicture intento {intento + 1} fallido: {e}")
            time.sleep(1 + intento)

    logging.error(f"No se pudo capturar imagen: {ruta_png}")
    return False


def generar_imagenes(trabajos: list[tuple]) -> dict:
    import xlwings as xw
    import subprocess

    resultados = {ruta_png: False for _, _, ruta_png in trabajos}
    xlsx_temps = []

    for df, titulo, ruta_png in trabajos:
        ruta_xlsx = ruta_png.replace(".png", ".xlsx")
        _generar_xlsx(df, titulo, ruta_xlsx)
        xlsx_temps.append((ruta_xlsx, ruta_png, len(df) + 2, len(df.columns)))

    app = None
    try:
        subprocess.run(["taskkill", "/f", "/im", "EXCEL.EXE"], capture_output=True)
        time.sleep(1)
        app = xw.App(visible=True, add_book=False)
        app.display_alerts = False

        for ruta_xlsx, ruta_png, n_filas, n_cols in xlsx_temps:
            if not os.path.exists(ruta_xlsx):
                continue
            wb_xw = None
            try:
                wb_xw = app.books.open(os.path.normpath(ruta_xlsx))
                ws_xw = wb_xw.sheets["Tabla"]
                ws_xw.activate()
                time.sleep(0.5)
                resultados[ruta_png] = _capturar_xlsx_a_png(ws_xw, n_filas, n_cols, ruta_png)
            except Exception as e:
                logging.error(f"Error capturando {ruta_png}: {e}")
            finally:
                if wb_xw:
                    try:
                        wb_xw.close()
                    except Exception:
                        pass
                try:
                    os.remove(ruta_xlsx)
                except Exception:
                    pass
                time.sleep(0.5)
    except Exception as e:
        logging.error(f"Error iniciando Excel para capturas: {e}")
    finally:
        if app:
            try:
                app.quit()
            except Exception:
                pass

    return resultados


# ── Mensaje caption ───────────────────────────────────────────────────────────

def _construir_caption(df_zonal: pd.DataFrame, fecha: date) -> str:
    fecha_str = fecha.strftime("%d/%m/%Y")
    fila_total = df_zonal[df_zonal["ZONAL"] == "TOTAL"]
    if fila_total.empty:
        return f"Asistencia PLANILLA - {fecha_str}"
    planilla   = int(fila_total.iloc[0]["PLANILLA"])
    asistentes = int(fila_total.iloc[0]["ASISTENTES"])
    pct        = float(fila_total.iloc[0]["%ASISTENCIA"])
    emoji = "✅" if pct >= 90 else ("⚠️" if pct >= 70 else "🔴")
    return (
        f"Buenos días,\n"
        f"*Asistencia PLANILLA — {fecha_str}*\n"
        f"{emoji} *{asistentes} / {planilla}* vendedores presentes ({pct:.1f}%)"
    )


# ── Flujo principal ───────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Informe de asistencia PLANILLA")
    parser.add_argument("--fecha",   default=None, help="Fecha YYYY-MM-DD (default: hoy)")
    parser.add_argument("--destino", default=None, help="Destino WA en lugar del grupo (ID o nombre de config.json)")
    args = parser.parse_args()

    fecha_hoy = (
        datetime.strptime(args.fecha, "%Y-%m-%d").date()
        if args.fecha else date.today()
    )
    destino = args.destino if args.destino else GRUPO_GESTION

    logging.info(f"=== Informe asistencia PLANILLA — {fecha_hoy} | destino: {destino} ===")

    # 1. Autenticar y cargar datos
    service = _autenticar_sheets()

    df_rh = cargar_rh(service)
    if df_rh.empty:
        logging.error("RH vacio — abortando.")
        sys.exit(1)

    df_asistencia = cargar_asistencia(service, fecha_hoy)
    if df_asistencia.empty:
        logging.warning(f"Sin registros de asistencia para {fecha_hoy} — se continua con 0 asistentes.")

    # 2. Calcular tablas
    df_zonal, df_supervisor = calcular_tablas(df_rh, df_asistencia)
    if df_zonal.empty:
        logging.error("No se pudo construir tabla por zonal — abortando.")
        sys.exit(1)

    logging.info("Tabla por ZONAL:\n" + df_zonal.to_string(index=False))
    if not df_supervisor.empty:
        logging.info("Tabla por SUPERVISOR:\n" + df_supervisor.to_string(index=False))

    # 3. Generar imágenes
    fecha_str = fecha_hoy.strftime("%d/%m/%Y")
    png_zonal = str(TEMP_DIR / f"asistencia_{fecha_hoy}_zonal.png")
    png_sup   = str(TEMP_DIR / f"asistencia_{fecha_hoy}_sup.png")

    trabajos = [(df_zonal, f"ASISTENCIA PLANILLA POR ZONAL — {fecha_str}", png_zonal)]
    if not df_supervisor.empty:
        trabajos.append((df_supervisor, f"ASISTENCIA PLANILLA POR SUPERVISOR — {fecha_str}", png_sup))

    resultados = generar_imagenes(trabajos)
    ok_zonal = resultados.get(png_zonal, False)
    ok_sup   = resultados.get(png_sup,   False)

    # 4. Enviar por WhatsApp
    wa      = WhatsAppClient()
    caption = _construir_caption(df_zonal, fecha_hoy)

    # IDs de menciones: leer desde contacts del config
    _cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    _contacts = _cfg.get("contacts", {})
    menciones_ids = [
        _contacts[k]
        for k in ("Jesús Ascencios", "Carlos P")
        if k in _contacts
    ]
    # El texto debe tener @número (sin @c.us), no @nombre — así WhatsApp renderiza la mención real
    partes_mencion = " ".join(
        f"@{wa_id.replace('@c.us', '')}" for wa_id in menciones_ids
    )
    texto_menciones = caption + "\n" + partes_mencion

    logging.info(f"Enviando informe a: {destino}")
    if ok_zonal:
        wa.send_image(destino, png_zonal, caption=caption)
        time.sleep(3)
        wa.send_mention(destino, texto_menciones, menciones_ids)
    else:
        logging.warning("Sin imagen zonal — enviando solo texto con menciones.")
        wa.send_mention(destino, texto_menciones, menciones_ids)

    if not df_supervisor.empty:
        time.sleep(8)
        if ok_sup:
            wa.send_image(destino, png_sup, caption="")
        else:
            logging.warning("Sin imagen de supervisores.")

    # 5. Enviar lista de ausentes
    time.sleep(5)
    texto_ausentes = calcular_ausentes(df_rh, df_asistencia)
    if texto_ausentes:
        wa.send_text(destino, texto_ausentes)
        logging.info("Lista de ausentes enviada.")

    # Limpiar temporales
    for f in [png_zonal, png_sup]:
        try:
            if os.path.exists(f):
                os.remove(f)
        except Exception:
            pass

    logging.info("=== Informe asistencia completado ===")


if __name__ == "__main__":
    main()
